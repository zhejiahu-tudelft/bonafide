"""Material financial, cutoff, option-method and deliverable checks."""
from pathlib import Path
import sys, json, math, hashlib, base64, io, re
from urllib.parse import unquote
import numpy as np
import pandas as pd
import openpyxl
from bs4 import BeautifulSoup
from PIL import Image
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
from common.code.workbook_validation import validate_formula_caches
from common.code.options_tools import bsm_greeks
from common.code.options_strategy_tools import aggregate,expiry_payoff,execution_price
from config import ROOT,CUTOFF,PRICE,price_history
from valuation import SCENARIOS,terminal_fcff,SHARES
D=ROOT/'data/processed_data'

def build():
    checks=[]
    def passed(name,details): checks.append(dict(check=name,status='passed',details=details))
    h=pd.read_csv(D/'financial_history.csv',index_col='period')
    assert list(h.index)==['2021','2022','2023','2024','2025','H1 2025','H1 2026','TTM']
    for row in h.itertuples():
        assert math.isclose(row.assets,row.liabilities+row.equity,abs_tol=1e-6)
        assert math.isclose(row.cfo-row.ppe_capex-row.mine_capex,row.fcf,abs_tol=1e-6)
        assert math.isclose(row.shares_issued-row.treasury_shares,row.actual_shares,abs_tol=1e-9)
        assert math.isclose(row.debt-row.cash,row.net_debt,abs_tol=1e-6)
    assert math.isclose(h.loc['2023','mine_capex'],9.326)
    assert math.isclose(h.loc['2023','fcf'],24.538,abs_tol=1e-6)
    assert math.isclose(h.loc[['2021','2022','2023','2024','2025'],'fcf'].sum(),80.904,abs_tol=1e-6)
    assert math.isclose(h.loc['TTM','fcf'],32.439,abs_tol=1e-6)
    for field in ['revenue','va_sales','net_income','cfo','capex','da','sbc']:
        assert math.isclose(h.loc['TTM',field],h.loc['2025',field]+h.loc['H1 2026',field]-h.loc['H1 2025',field],abs_tol=1e-6)
    assert (pd.Timestamp(h.loc['TTM','end'])-pd.Timestamp(h.loc['TTM','start'])).days+1==371
    assert abs(h.loc['TTM','net_income']/h.loc['TTM','diluted_shares']-h.loc['TTM','eps'])<.01
    passed('Financial identities and TTM','Eight periods; corrected 2023 mine capex; 371-day trailing calculation; cumulative FCF $80.904m.')
    # Independently verify the custom mining values against the actual filing tables.
    filing=ROOT/'report/company_filings/10-K_2026-02-12_mtrn-20251231.htm'
    source=BeautifulSoup(filing.read_text(),'html.parser')
    cash_rows=[x.get_text(' ',strip=True) for x in source.find_all('tr') if 'Payments for mine development' in x.get_text(' ',strip=True)]
    assert cash_rows, 'Mine development line missing from primary filing'
    flattened=' '.join(cash_rows)
    for figure in ['26,288','12,159','9,326']: assert figure in flattened
    passed('Primary filing cash-flow check','2023–2025 mine spending matched to the retained 2025 10-K table.')
    for filename in ['financial_audit.csv','peer_financial_audit.csv']:
        a=pd.read_csv(D/filename)
        dates=a.filed.dropna().astype(str)
        assert (dates<=CUTOFF).all(),filename
    for f in (ROOT/'data/market_data').glob('*_daily.csv'):
        p=pd.read_csv(f,index_col='date',parse_dates=True)
        assert str(p.index[-1].date())<=CUTOFF
        assert p.index.is_unique and p.index.is_monotonic_increasing
        if f.name!='CAD=X_daily.csv': assert str(p.index[-1].date())==CUTOFF
    assert math.isclose(price_history('MTRN').close.iloc[-1],PRICE,abs_tol=.005)
    passed('Information cutoff','Selected SEC filings and all market histories respect September 18, 2026; close reconciles to $251.53.')
    v=json.loads((D/'valuation_summary.json').read_text());c=v['claims']
    claims=c['debt_principal']+c['finance_leases']+c['retirement']-c['cash']
    assert math.isclose(claims,v['net_claims_dcf'])
    assert math.isclose(v['market_cap'],20.834*PRICE)
    assert math.isclose(v['ev_standard'],v['market_cap']+claims-c['retirement'])
    summary=pd.read_csv(D/'dcf_summary.csv').set_index('scenario')
    for name,cfg in SCENARIOS.items():
        frame=pd.read_csv(D/f'dcf_{name.lower()}.csv');row=summary.loc[name]
        assert (pd.to_datetime(frame.period_end)>pd.Timestamp(CUTOFF)).all()
        np.testing.assert_allclose(frame.fcff,frame.ebit*(1-frame.tax_rate)+frame.da-frame.capex-frame.change_nwc)
        terminal=terminal_fcff(frame,cfg)
        explicit=sum(cf/(1+cfg['wacc'])**(i+1) for i,cf in enumerate(frame.fcff))
        ev=explicit+terminal/(cfg['wacc']-cfg['g'])/(1+cfg['wacc'])**5
        assert math.isclose(ev,row.enterprise)
        assert math.isclose((ev-claims)/SHARES,row.per_share)
    passed('DCF and enterprise-equity bridge','All scenarios independently recalculated; terminal growth resets working-capital reinvestment.')
    wbpath=ROOT/'excel/MTRN_Research_Model.xlsx'
    formula=openpyxl.load_workbook(wbpath,data_only=False)
    cached=openpyxl.load_workbook(wbpath,data_only=True)
    count=validate_formula_caches(formula,cached)
    for name in SCENARIOS:
        assert math.isclose(cached['DCF '+name]['B38'].value,summary.loc[name,'per_share'])
    passed('Excel formula evaluation',f'{count} formulas independently evaluated from their precedents and matched to cached results across {len(formula.sheetnames)} sheets.')
    a=pd.read_csv(D/'options_audit.csv');raw=json.loads((ROOT/'data/options_data/cboe_raw.json').read_text())
    assert raw['data']['last_trade_time'][:10]==CUTOFF
    assert raw['data']['current_price']==PRICE
    assert len(a)==352 and a.usable.sum()==0
    assert a.contract.is_unique and a.multiplier.eq(100).all()
    assert (a.expiration==pd.to_datetime(a.contract.str.extract(r'(\d{6})[CP]\d{8}$')[0],format='%y%m%d').dt.strftime('%Y-%m-%d')).all()
    assert np.allclose(a.strike,a.contract.str[-8:].astype(int)/1000)
    unexpired=a[a.days_to_expiration>0]
    assert len(unexpired)==254 and unexpired.open_interest.sum()==755 and unexpired.volume.sum()==18
    passed('Options snapshot quality','Real 352-record chain; OCC dates/strikes and 100-share multiplier verified; 254 unexpired records; zero eligible strategy legs.')
    # Generic model diagnostics use a synthetic unit-test input, not MTRN quotes.
    spot,strike,time,rate,div,sigma=100.,105.,.7,.04,.01,.30
    call=bsm_greeks(spot,strike,time,rate,div,sigma,'call');put=bsm_greeks(spot,strike,time,rate,div,sigma,'put')
    assert math.isclose(call['price']-put['price'],spot*math.exp(-div*time)-strike*math.exp(-rate*time),abs_tol=1e-10)
    step=1e-4
    up=bsm_greeks(spot,strike,time,rate,div,sigma+step,'call');down=bsm_greeks(spot,strike,time,rate,div,sigma-step,'call')
    assert math.isclose(call['vanna'],(up['delta']-down['delta'])/(2*step),abs_tol=1e-6)
    assert math.isclose(call['vomma'],(up['price']-2*call['price']+down['price'])/step**2,abs_tol=1e-4)
    legs=[dict(call,quantity=1,multiplier=100,strike=105,type='call'),dict(put,quantity=-1,multiplier=100,strike=105,type='put')]
    result=aggregate(legs)
    assert math.isclose(result['delta'],100*math.exp(-div*time))
    assert abs(result['gamma'])<1e-8 and abs(result['vega'])<1e-8
    np.testing.assert_allclose(expiry_payoff(legs,[80,105,130]),100*(np.array([80,105,130])-105)-result['premium_paid'])
    assert execution_price(2,3,1)==3 and execution_price(2,3,-1)==2
    passed('Reusable options conventions','Put-call parity, vanna/vomma finite differences, signed contract aggregation, payoff and bid/ask execution-side checks passed. Synthetic checks are not research quotes.')
    report=ROOT/'final/MTRN_Investment_Research.html';html=report.read_text();soup=BeautifulSoup(html,'html.parser')
    assert not re.search(r'\{\{[A-Za-z_0-9]+\}\}',html)
    assert len(soup.select('main h2'))==18 and len(soup.select('main img'))==6
    for img in soup.select('main img'):
        assert img['alt']
        Image.open(io.BytesIO(base64.b64decode(img['src'].split(',',1)[1]))).verify()
    ids={x['id'] for x in soup.select('[id]')}
    for link in soup.select('a[href]'):
        target=unquote(link['href'])
        if target.startswith('#'): assert target[1:] in ids,target
        elif not target.startswith(('http:','https:','mailto:')): assert (report.parent/target).resolve().exists(),target
    d=json.loads(soup.find(id='dcf-inputs').string)
    w=.105;g=.03
    terminal=d['lastVA']*(1+g)*((d['margin']-d['da'])*(1-d['tax'])+d['da']-d['capex'])-d['lastVA']*g*d['nwc']
    value=(sum(cf/(1+w)**(i+1) for i,cf in enumerate(d['cashflows']))+terminal/(w-g)/(1+w)**5-d['claims'])/d['shares']
    assert math.isclose(value,summary.loc['Base','per_share'])
    passed('Final artifact integrity','18 sections; six valid embedded charts; all local and section links exist; no unresolved template fields; interactive base DCF matches Python.')
    inventory=pd.read_csv(D/'evidence_inventory.csv')
    for r in inventory.itertuples():
        p=ROOT/r.local_path
        assert hashlib.sha256(p.read_bytes()).hexdigest()==r.sha256,r.local_path
    sources=pd.read_csv(ROOT/'report/source_log.csv')
    for path in sources.local_path: assert (ROOT/path).exists(),path
    passed('Source archive integrity',f'{len(inventory)} file hashes verified; {len(sources)} source entries point to retained evidence.')
    result=dict(status='passed',cutoff=CUTOFF,checks=checks)
    (D/'validation_results.json').write_text(json.dumps(result,indent=2))
    print(json.dumps(result,indent=2))

if __name__=='__main__': build()
